import openpyxl
import json
import sys

def read_excel_formulas(file_path):
    try:
        # Load the workbook with data_only=False to read formulas
        wb = openpyxl.load_workbook(file_path, data_only=False)
        
        # Also load with data_only=True to get the evaluated values
        wb_data = openpyxl.load_workbook(file_path, data_only=True)
        
        results = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            ws_data = wb_data[sheet_name]
            
            sheet_data = []
            
            # Read first 30 rows and 15 columns for preview
            for row_idx in range(1, 31):
                row_info = []
                has_content = False
                for col_idx in range(1, 16):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell_data = ws_data.cell(row=row_idx, column=col_idx)
                    
                    val = cell.value
                    if val is not None:
                        has_content = True
                        
                    is_formula = isinstance(val, str) and val.startswith('=')
                    
                    cell_info = {
                        'coord': cell.coordinate,
                        'value': cell_data.value,
                        'formula': val if is_formula else None
                    }
                    row_info.append(cell_info)
                
                if has_content:
                    sheet_data.append(row_info)
            
            if sheet_data:
                results[sheet_name] = sheet_data
                
        # Print a readable summary
        for sheet, rows in results.items():
            print(f"\n--- Hoja: {sheet} ---")
            for r_idx, row in enumerate(rows):
                # Print only cells with data
                cells_with_data = [c for c in row if c['value'] is not None or c['formula'] is not None]
                if cells_with_data:
                    row_str = " | ".join([f"{c['coord']}: {c['formula'] if c['formula'] else c['value']}" for c in cells_with_data])
                    print(row_str)
                    
    except Exception as e:
        print(f"Error: {e}")

if __name__ == '__main__':
    read_excel_formulas(r"C:\Users\keine\.gemini\antigravity\scratch\checkadmin\2026 FORMULAS Y COSTOS LM.xlsx")
